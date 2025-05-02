import tkinter as tk
from openpyxl import Workbook, load_workbook
import os

if not os.path.exists("student_scores.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(["Student Name", "Score", "Remarks"])
    wb.save("student_scores.xlsx")

window = tk.Tk()
window.title("Score Tracker")

tk.Label(window, text="Student Name").grid(row=0, column=0)
tk.Label(window, text="Score").grid(row=1, column=0)

name_entry = tk.Entry(window)
score_entry = tk.Entry(window)
name_entry.grid(row=0, column=1)
score_entry.grid(row=1, column=1)

output_label = tk.Label(window, text="", fg="blue")
output_label.grid(row=4, column=0, columnspan=2)

def submit():
    name = name_entry.get()
    try:
        score = int(score_entry.get())
        if score < 0 or score > 100:
            remark = "Invalid"
        elif score > 75:
            remark = "Passed!"
        else:
            remark = "Failed"
    except:
        remark = "Invalid"

    wb = load_workbook("student_scores.xlsx")
    ws = wb.active

    if ws.cell(row=ws.max_row, column=1).value == "Average":
        ws.delete_rows(ws.max_row)

    if remark != "Invalid":
        ws.append([name, score, remark])
    else:
        output_label.config(text="Score Invalid!")
        name_entry.delete(0, tk.END)
        score_entry.delete(0, tk.END)

    total = 0
    count = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[2] in ["Passed!", "Failed"]:
            total += row[1]
            count += 1
    if count > 0:
        avg = total / count
        ws.append(["Average", avg, ""])

    wb.save("student_scores.xlsx")
    if remark != 'Invalid':
        output_label.config(text="Score submitted!")

    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

def view_scores():
    wb = load_workbook("student_scores.xlsx")
    ws = wb.active

    records_window = tk.Toplevel(window)
    records_window.title("Student Records")

    tk.Label(records_window, text="Name", font=('Times New Roman', 10, 'bold')).grid(row=0, column=0, padx=10, pady=5)
    tk.Label(records_window, text="Score", font=('Times New Roman', 10, 'bold')).grid(row=0, column=1, padx=10, pady=5)
    tk.Label(records_window, text="Remarks", font=('Times New Roman', 10, 'bold')).grid(row=0, column=2, padx=10, pady=5)

    row_num = 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        tk.Label(records_window, text=row[0]).grid(row=row_num, column=0, padx=10)
        tk.Label(records_window, text=row[1]).grid(row=row_num, column=1, padx=10)
        tk.Label(records_window, text=row[2]).grid(row=row_num, column=2, padx=10)
        row_num += 1

submit_btn = tk.Button(window, text="Submit", width=12, command=submit)
submit_btn.grid(row=2, column=0, padx=10, pady=5)

view_btn = tk.Button(window, text="View Scores", width=12, command=view_scores)
view_btn.grid(row=2, column=1, padx=5, pady=5)

window.mainloop()
